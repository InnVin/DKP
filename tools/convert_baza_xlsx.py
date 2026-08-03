from __future__ import annotations

from copy import copy
from pathlib import Path

import xlrd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "MyFiles" / "BAZA.xls"
TARGET = ROOT / "MyFiles" / "BAZA.xlsx"


def _color(book: xlrd.Book, index: int | None, fallback: str = "000000") -> str:
    if index is None or index in (64, 65, 32767):
        return fallback
    rgb = book.colour_map.get(index)
    return "".join(f"{part:02X}" for part in rgb) if rgb else fallback


def _side(book: xlrd.Book, line_style: int, color_index: int) -> Side:
    styles = {
        0: None,
        1: "thin",
        2: "medium",
        3: "dashed",
        4: "dotted",
        5: "thick",
        6: "double",
        7: "hair",
        8: "mediumDashed",
        9: "dashDot",
        10: "mediumDashDot",
        11: "dashDotDot",
        12: "mediumDashDotDot",
        13: "slantDashDot",
    }
    style = styles.get(line_style)
    return Side(style=style, color=_color(book, color_index)) if style else Side()


def convert() -> Path:
    source = xlrd.open_workbook(str(SOURCE), formatting_info=True)
    original = source.sheet_by_name("Пуст")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Пуст"
    sheet.sheet_view.showGridLines = False

    horizontal = {0: "general", 1: "left", 2: "center", 3: "right", 4: "fill", 5: "justify", 6: "centerContinuous", 7: "distributed"}
    vertical = {0: "top", 1: "center", 2: "bottom", 3: "justify", 4: "distributed"}

    for row in range(original.nrows):
        row_info = original.rowinfo_map.get(row)
        if row_info:
            sheet.row_dimensions[row + 1].height = row_info.height / 20
            sheet.row_dimensions[row + 1].hidden = bool(row_info.hidden)
        for column in range(original.ncols):
            source_cell = original.cell(row, column)
            target_cell = sheet.cell(row + 1, column + 1, source_cell.value if source_cell.value != "" else None)
            xf = source.xf_list[source_cell.xf_index]
            source_font = source.font_list[xf.font_index]
            target_cell.font = Font(
                name=source_font.name or "Times New Roman",
                size=(source_font.height or 200) / 20,
                bold=bool(source_font.bold),
                italic=bool(source_font.italic),
                underline="single" if source_font.underlined else None,
                color=_color(source, source_font.colour_index),
            )
            target_cell.alignment = Alignment(
                horizontal=horizontal.get(xf.alignment.hor_align, "general"),
                vertical=vertical.get(xf.alignment.vert_align, "bottom"),
                wrap_text=bool(xf.alignment.text_wrapped),
                shrink_to_fit=bool(xf.alignment.shrink_to_fit),
                text_rotation=xf.alignment.rotation,
            )
            target_cell.border = Border(
                left=_side(source, xf.border.left_line_style, xf.border.left_colour_index),
                right=_side(source, xf.border.right_line_style, xf.border.right_colour_index),
                top=_side(source, xf.border.top_line_style, xf.border.top_colour_index),
                bottom=_side(source, xf.border.bottom_line_style, xf.border.bottom_colour_index),
            )
            if xf.background.fill_pattern:
                target_cell.fill = PatternFill(
                    "solid",
                    fgColor=_color(source, xf.background.pattern_colour_index, "FFFFFF"),
                )
            target_cell.number_format = source.format_map[xf.format_key].format_str

    for column, info in original.colinfo_map.items():
        dimension = sheet.column_dimensions[get_column_letter(column + 1)]
        dimension.width = max(1, info.width / 256)
        dimension.hidden = bool(info.hidden)

    for row_start, row_end, column_start, column_end in original.merged_cells:
        sheet.merge_cells(
            start_row=row_start + 1,
            end_row=row_end,
            start_column=column_start + 1,
            end_column=column_end,
        )

    gray = PatternFill("solid", fgColor="E7E9E8")
    for address in (
        "A4:C4", "J4", "B9:J9", "D10:F10", "H10:J10", "D11:J11", "D12:J12",
        "B15:J15", "D16:F16", "H16:J16", "D17:J17", "D18:J18",
        "D22:J31", "D34:F34", "D43:E43", "G43:J43", "D46:E46", "G46:J46",
    ):
        target = sheet[address]
        if hasattr(target, "coordinate"):
            target.fill = copy(gray)
            continue
        for row in target:
            for cell in row:
                cell.fill = copy(gray)

    sheet["A4"].alignment = copy(sheet["A4"].alignment)
    sheet["A4"].alignment = Alignment(horizontal="center", vertical="center")
    sheet["J4"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = None
    sheet.print_area = "A1:J47"
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.page_margins.left = 0.2
    sheet.page_margins.right = 0.2
    sheet.page_margins.top = 0.25
    sheet.page_margins.bottom = 0.25
    sheet.page_margins.header = 0
    sheet.page_margins.footer = 0
    sheet.print_options.horizontalCentered = True
    workbook.save(TARGET)
    return TARGET


if __name__ == "__main__":
    print(convert())
