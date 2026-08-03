import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image
import pypdfium2 as pdfium

from app import database
from app.excel_service import _prepare_baza_payload, _write_default_workbook, create_contract
from app.field_postprocess import postprocess_fields
from app.hybrid_ocr import _resolve_candidates
from app.ocr import extract_fields
from app.main import _friendly_ocr_error


ROOT = Path(__file__).resolve().parents[1]


class WebPwaUiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        static = ROOT / "app" / "static"
        cls.index = (static / "index.html").read_text(encoding="utf-8")
        cls.styles = (static / "styles.css").read_text(encoding="utf-8")
        cls.script = (static / "pwa.js").read_text(encoding="utf-8")
        cls.ui_config = (static / "ui-config.js").read_text(encoding="utf-8")

    def test_three_steps_and_mobile_navigation(self):
        for tab in ("documents", "data", "review"):
            self.assertIn(f'data-deal-tab="{tab}"', self.index)
        self.assertIn('class="bottom-navigation"', self.index)
        self.assertIn('{ action: "deal", label: "ДКП" }', self.ui_config)
        self.assertIn('{ action: "archive", label: "Архив" }', self.ui_config)
        self.assertNotIn('{ action: "settings", label: "Настройки" }', self.ui_config)
        self.assertNotIn("Профиль", self.ui_config)

    def test_document_actions_and_multiple_upload(self):
        self.assertIn("<h1>Загрузите документы</h1>", self.index)
        self.assertIn('id="reprocessDeal"', self.index)
        self.assertGreaterEqual(self.index.count("multiple>"), 5)
        self.assertNotIn(">Повернуть<", self.index)
        self.assertIn("camera-symbol", self.index)
        self.assertIn("scan-symbol", self.index)

    def test_accordions_themes_and_fixed_action(self):
        for section in ("contract", "seller", "buyer", "vehicle", "notes"):
            self.assertIn(f'id="section-{section}"', self.index)
        self.assertIn('id="themeToggle"', self.index)
        self.assertNotIn('data-theme-choice="system"', self.index)
        self.assertIn(".floating-next", self.styles)
        self.assertIn("position: fixed", self.styles)
        self.assertIn("env(safe-area-inset-bottom)", self.styles)

    def test_swipe_toast_and_double_back_logic(self):
        self.assertIn('addEventListener("touchstart"', self.script)
        self.assertIn('addEventListener("touchend"', self.script)
        self.assertIn("Нажмите «Назад» ещё раз", self.script)
        self.assertIn("Отменить", self.script)
        self.assertIn("2000", self.script)

    def test_ready_files_and_archive(self):
        self.assertIn("Готовые файлы", self.index)
        self.assertIn("navigator.share", self.script)
        self.assertIn('window.open(`${file.download_url}?inline=1`', self.script)
        self.assertIn("download_url", self.script)
        self.assertIn("archive-item", self.script)
        self.assertIn("data-toggle-archive", self.script)
        self.assertIn("data-archive-files", self.script)
        self.assertIn("data-rename-file", self.script)


class OcrExtractionTests(unittest.TestCase):
    def test_ocr_connection_error_is_user_friendly(self):
        message = _friendly_ocr_error(RuntimeError("<urlopen error [WinError 10013] blocked>"))
        self.assertIn("нет соединения", message)
        self.assertNotIn("WinError", message)

    def test_conflicting_candidates_require_user_choice(self):
        resolved, conflicts = _resolve_candidates(
            {
                "body_number": [
                    {"value": "SCP92-1009017", "page": 1, "confidence": 0.91},
                    {"value": "8CP92-1009017", "page": 2, "confidence": 0.89},
                ]
            },
            "vehicle_docs",
        )
        self.assertNotIn("body_number", resolved)
        self.assertEqual(len(conflicts["body_number"]), 2)

    def test_candidate_majority_is_resolved(self):
        resolved, conflicts = _resolve_candidates(
            {
                "body_number": [
                    {"value": "SCP92-1009017", "page": 1, "confidence": 0.91},
                    {"value": "SCP92-1009017", "page": 2, "confidence": 0.88},
                    {"value": "8CP92-1009017", "page": 3, "confidence": 0.90},
                ]
            },
            "vehicle_docs",
        )
        self.assertEqual(resolved["body_number"], "SCP92-1009017")
        self.assertNotIn("body_number", conflicts)

    def test_body_number_has_no_model_specific_replacement(self):
        result = postprocess_fields({"body_number": "8CP92-1009017"}, "vehicle_docs")
        self.assertEqual(result["body_number"], "8CP92-1009017")

    def test_extract_vehicle_fields(self):
        text = """
        VIN: XTA210990Y1234567
        Марка, модель: LADA 21099
        Год выпуска: 2000
        Цвет кузова: СЕРЕБРИСТЫЙ
        """
        fields = extract_fields(text, "sts")
        self.assertEqual(fields["vin"], "XTA210990Y1234567")
        self.assertEqual(fields["vehicle_year"], "2000")


class GeneratedFileTests(unittest.TestCase):
    def test_generated_file_rename_is_scoped_and_keeps_extension(self):
        originals = {
            name: getattr(database, name)
            for name in ("DATA_DIR", "DEALS_DIR", "UPLOAD_DIR", "OUTPUT_DIR", "LEGACY_OUTPUT_DIR", "DB_PATH")
        }
        try:
            with tempfile.TemporaryDirectory() as folder:
                root = Path(folder)
                database.DATA_DIR = root
                database.DEALS_DIR = root / "deals"
                database.UPLOAD_DIR = root / "uploads"
                database.OUTPUT_DIR = root / "contracts"
                database.LEGACY_OUTPUT_DIR = root / "legacy"
                database.DB_PATH = root / "test.sqlite3"
                database.init_db()
                deal_id = database.save_deal({"seller_full_name": "Тест"})
                source = database.generated_dir(deal_id) / "Договор.pdf"
                source.write_bytes(b"%PDF-test")
                stored = database.replace_generated_files(deal_id, [{
                    "kind": "pdf",
                    "name": source.name,
                    "stored_path": str(source),
                    "media_type": "application/pdf",
                }])[0]
                renamed = database.rename_generated_file(deal_id, stored["id"], "Новый договор.pdf")
                self.assertEqual(renamed["name"], "Новый договор.pdf")
                self.assertTrue(Path(renamed["stored_path"]).exists())
                with self.assertRaises(ValueError):
                    database.rename_generated_file(deal_id, stored["id"], "Новый договор.exe")
        finally:
            for name, value in originals.items():
                setattr(database, name, value)


class ExcelTests(unittest.TestCase):
    def test_baza_date_format(self):
        payload = _prepare_baza_payload({"contract_date": "2026-06-18"})
        self.assertEqual(payload["contract_date"], "18.06.2026")

    def test_canonical_baza_outputs_xlsx_pdf_jpg(self):
        payload = {
            "contract_place": "г. Якутск",
            "contract_date": "2026-07-30",
            "seller_full_name": "Копылов Анатолий Петрович",
            "seller_passport": "9821 988057",
            "seller_passport_issue_date": "2022-03-14",
            "seller_passport_issued_by": "МВД ПО РЕСПУБЛИКЕ САХА (ЯКУТИЯ)",
            "seller_address": "РЕСП. САХА /ЯКУТИЯ/",
            "buyer_full_name": "Васильева Анастасия Дмитриевна",
            "buyer_passport": "9822 040066",
            "buyer_passport_issue_date": "2023-04-14",
            "buyer_passport_issued_by": "МВД ПО РЕСПУБЛИКЕ САХА (ЯКУТИЯ)",
            "buyer_address": "РЕСП. САХА /ЯКУТИЯ/",
            "vehicle_make_model": "TOYOTA BELTA",
            "vehicle_type": "Легковой седан",
            "vehicle_year": "2006",
            "vin": "ОТСУТСТВУЕТ",
            "body_number": "SCP92-1009017",
            "chassis_number": "ОТСУТСТВУЕТ",
            "color": "Серый",
            "registration_plate": "В831ЕТ14",
            "pts_series_number": "14РС 615017",
            "sts_series_number": "14РС 615017",
            "price": "750000",
        }
        original_deals_dir = database.DEALS_DIR
        try:
            with tempfile.TemporaryDirectory() as folder:
                database.DEALS_DIR = Path(folder)
                files = create_contract(77, payload)
                self.assertEqual({item["kind"] for item in files}, {"xlsx", "pdf", "jpg"})
                paths = {item["kind"]: Path(item["stored_path"]) for item in files}
                self.assertTrue(all(path.exists() and path.stat().st_size > 1000 for path in paths.values()))
                workbook = load_workbook(paths["xlsx"])
                sheet = workbook["Пуст"]
                self.assertEqual(sheet["D22"].value, "TOYOTA BELTA")
                self.assertEqual(sheet["D22"].fill.fgColor.rgb[-6:], "E7E9E8")
                self.assertEqual(sheet["A4"].alignment.horizontal, "center")
                self.assertEqual(sheet.print_area, "'Пуст'!$A$1:$J$47")
                self.assertEqual(sheet.page_setup.fitToWidth, 1)
                self.assertEqual(sheet.page_setup.fitToHeight, 1)
                workbook.close()
                self.assertEqual(paths["pdf"].read_bytes()[:4], b"%PDF")
                pdf = pdfium.PdfDocument(str(paths["pdf"]))
                self.assertEqual(len(pdf), 1)
                pdf.close()
                with Image.open(paths["jpg"]) as image:
                    self.assertTrue(1238 <= image.width <= 1242)
                    self.assertTrue(1752 <= image.height <= 1756)
        finally:
            database.DEALS_DIR = original_deals_dir

    def test_default_workbook_created(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "contract.xlsx"
            seller_name = "Петров Пётр Петрович"
            buyer_name = "Иванов Иван Иванович"
            _write_default_workbook(
                path,
                {
                    "buyer_full_name": buyer_name,
                    "seller_full_name": seller_name,
                    "vin": "XTA210990Y1234567",
                },
            )
            workbook = load_workbook(path, read_only=True, data_only=False)
            values = [str(cell.value or "") for row in workbook.active.iter_rows() for cell in row]
            self.assertGreaterEqual(sum(seller_name in value for value in values), 2)
            self.assertGreaterEqual(sum(buyer_name in value for value in values), 2)
            workbook.close()


if __name__ == "__main__":
    unittest.main()
