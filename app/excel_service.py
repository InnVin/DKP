from __future__ import annotations

import re
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any

import pypdfium2 as pdfium
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
import xlrd
from xlutils.copy import copy as copy_xls

from . import database


ROOT = database.ROOT
TEMPLATE_PATH = ROOT / "MyFiles" / "BAZA.xls"
FONT_REGULAR = Path(r"C:\Windows\Fonts\arial.ttf")
FONT_BOLD = Path(r"C:\Windows\Fonts\arialbd.ttf")
INPUT_FILL = PatternFill("solid", fgColor="E7E9E8")
BLACK_SIDE = Side(style="thin", color="000000")

CELL_MAP = {
    "A4": "contract_place",
    "J4": "contract_date",
    "B9": "seller_full_name",
    "D10": "seller_passport",
    "H10": "seller_passport_issue_date",
    "D11": "seller_passport_issued_by",
    "D12": "seller_address",
    "B15": "buyer_full_name",
    "D16": "buyer_passport",
    "H16": "buyer_passport_issue_date",
    "D17": "buyer_passport_issued_by",
    "D18": "buyer_address",
    "D22": "vehicle_make_model",
    "D23": "vehicle_type",
    "D24": "vehicle_year",
    "D25": "vin",
    "D26": "body_number",
    "D27": "chassis_number",
    "D28": "color",
    "D29": "pts_series_number",
    "D30": "sts_series_number",
    "D31": "registration_plate",
    "D34": "price",
    "G43": "seller_full_name",
    "G46": "buyer_full_name",
}


def _fill_legacy_template(path: Path, payload: dict[str, Any]) -> None:
    source = xlrd.open_workbook(str(TEMPLATE_PATH), formatting_info=True)
    workbook = copy_xls(source)
    sheet = workbook.get_sheet(0)
    data = _prepare_baza_payload(payload)
    data.setdefault("seller_passport_issued_by", data.get("seller_passport_issued", ""))
    data.setdefault("buyer_passport_issued_by", data.get("buyer_passport_issued", ""))
    for address, key in CELL_MAP.items():
        row, column = _cell_coordinates(address)
        existing_xf = sheet._Worksheet__rows[row]._Row__cells[column].xf_idx
        sheet.write(row, column, str(data.get(key, "") or "").strip())
        sheet._Worksheet__rows[row]._Row__cells[column].xf_idx = existing_xf
    for row in (12, 18):
        for column in range(10):
            existing_xf = sheet._Worksheet__rows[row]._Row__cells[column].xf_idx
            sheet.write(row, column, "")
            sheet._Worksheet__rows[row]._Row__cells[column].xf_idx = existing_xf
    sheet.print_area = "$A$1:$J$47"
    sheet.fit_num_pages = 1
    sheet.fit_width_to_pages = 1
    sheet.fit_height_to_pages = 1
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(str(path))


def _cell_coordinates(address: str) -> tuple[int, int]:
    match = re.fullmatch(r"([A-Z]+)(\d+)", address)
    if not match:
        raise ValueError(address)
    column = 0
    for letter in match.group(1):
        column = column * 26 + ord(letter) - 64
    return int(match.group(2)) - 1, column - 1

INPUT_RANGES = (
    "A4:C4", "J4", "B9:J9", "D10:F10", "H10:J10", "D11:J11", "D12:J12",
    "B15:J15", "D16:F16", "H16:J16", "D17:J17", "D18:J18", "D22:J31",
    "D34:F34", "D43:E43", "G43:J43", "D46:E46", "G46:J46",
)


def _safe_name(value: str) -> str:
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', " ", value).strip()
    return re.sub(r"\s+", "_", value)[:60] or "договор"


def _output_path(deal_id: int, payload: dict[str, Any], suffix: str) -> Path:
    folder = database.generated_dir(deal_id)
    person = payload.get("buyer_full_name") or payload.get("seller_full_name") or f"сделка_{deal_id}"
    return folder / f"ДКП_{_safe_name(str(person))}{suffix}"


def _format_ru_date(value: Any) -> Any:
    text = str(value or "").strip()
    try:
        return datetime.strptime(text, "%Y-%m-%d").strftime("%d.%m.%Y")
    except ValueError:
        return value


def _prepare_baza_payload(payload: dict[str, Any]) -> dict[str, Any]:
    prepared = dict(payload)
    for key in ("contract_date", "seller_passport_issue_date", "buyer_passport_issue_date"):
        prepared[key] = _format_ru_date(prepared.get(key, ""))
    return prepared


def _style_input_ranges(sheet) -> None:
    for address in INPUT_RANGES:
        target = sheet[address]
        if hasattr(target, "coordinate"):
            target.fill = copy(INPUT_FILL)
            continue
        for row in target:
            for cell in row:
                cell.fill = copy(INPUT_FILL)


def _configure_print(sheet) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.print_area = "A1:J47"
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.page_margins.left = 0.2
    sheet.page_margins.right = 0.2
    sheet.page_margins.top = 0.25
    sheet.page_margins.bottom = 0.25
    sheet.page_margins.header = 0
    sheet.page_margins.footer = 0
    sheet.print_options.horizontalCentered = True


def _fill_template(path: Path, payload: dict[str, Any]) -> None:
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError("Не найден шаблон MyFiles/BAZA.xls")
    workbook = load_workbook(TEMPLATE_PATH)
    sheet = workbook["Пуст"]
    data = _prepare_baza_payload(payload)
    for address, key in CELL_MAP.items():
        value = data.get(key)
        sheet[address] = "" if value is None else str(value).strip()
    _style_input_ranges(sheet)
    sheet["A4"].alignment = Alignment(horizontal="center", vertical="center")
    sheet["J4"].alignment = Alignment(horizontal="center", vertical="center")
    _configure_print(sheet)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()


def _write_default_workbook(path: Path, payload: dict[str, Any]) -> None:
    """Compatibility entry point used by tests and local tools."""
    if TEMPLATE_PATH.exists() and TEMPLATE_PATH.suffix.lower() == ".xlsx":
        _fill_template(path, payload)
        return
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Пуст"
    for column in range(1, 11):
        sheet.column_dimensions[get_column_letter(column)].width = 11
    sheet.merge_cells("A1:J2")
    sheet["A1"] = "ДОГОВОР КУПЛИ-ПРОДАЖИ ТРАНСПОРТНОГО СРЕДСТВА"
    sheet["A1"].font = Font(name="Arial", size=14, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    row = 4
    for label, key in (("Продавец", "seller_full_name"), ("Покупатель", "buyer_full_name"), ("VIN", "vin")):
        sheet.cell(row, 1, label)
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
        sheet.cell(row, 2, str(payload.get(key, "")))
        for column in range(1, 11):
            sheet.cell(row, column).border = Border(bottom=BLACK_SIDE)
        row += 1
    sheet["G43"] = str(payload.get("seller_full_name", ""))
    sheet["G46"] = str(payload.get("buyer_full_name", ""))
    _configure_print(sheet)
    workbook.save(path)
    workbook.close()


def _register_pdf_fonts() -> None:
    if "AutoDogovor" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("AutoDogovor", str(FONT_REGULAR)))
    if "AutoDogovorBold" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("AutoDogovorBold", str(FONT_BOLD)))


def _fill_rgb(cell) -> tuple[float, float, float] | None:
    color = cell.fill.fgColor
    if cell.fill.fill_type != "solid" or color.type != "rgb" or not color.rgb:
        return None
    value = color.rgb[-6:]
    return tuple(int(value[index:index + 2], 16) / 255 for index in (0, 2, 4))


def _wrap_pdf_text(text: str, font_name: str, font_size: float, width: float) -> list[str]:
    words = str(text or "").split()
    if not words:
        return []
    lines: list[str] = []
    current = words[0]
    for word in words[1:]:
        candidate = f"{current} {word}"
        if pdfmetrics.stringWidth(candidate, font_name, font_size) <= width:
            current = candidate
        else:
            lines.append(current)
            current = word
    lines.append(current)
    return lines


def _render_workbook_pdf(xlsx_path: Path, pdf_path: Path) -> None:
    _register_pdf_fonts()
    workbook = load_workbook(xlsx_path, data_only=True)
    sheet = workbook["Пуст"]
    page_width, page_height = A4
    margin_x, margin_y = 24.0, 22.0
    column_units = [sheet.column_dimensions[get_column_letter(column)].width or 8.43 for column in range(1, 11)]
    row_units = [sheet.row_dimensions[row].height or 15 for row in range(1, 48)]
    width_scale = (page_width - margin_x * 2) / sum(column_units)
    height_scale = (page_height - margin_y * 2) / sum(row_units)
    x_positions = [margin_x]
    for width in column_units:
        x_positions.append(x_positions[-1] + width * width_scale)
    y_positions = [page_height - margin_y]
    for height in row_units:
        y_positions.append(y_positions[-1] - height * height_scale)

    merged_anchors: dict[tuple[int, int], tuple[int, int]] = {}
    covered: set[tuple[int, int]] = set()
    for merged in sheet.merged_cells.ranges:
        merged_anchors[(merged.min_row, merged.min_col)] = (merged.max_row, merged.max_col)
        for row in range(merged.min_row, merged.max_row + 1):
            for column in range(merged.min_col, merged.max_col + 1):
                if (row, column) != (merged.min_row, merged.min_col):
                    covered.add((row, column))

    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    document = canvas.Canvas(str(pdf_path), pagesize=A4, pageCompression=1)
    for row in range(1, 48):
        for column in range(1, 11):
            if (row, column) in covered:
                continue
            end_row, end_column = merged_anchors.get((row, column), (row, column))
            cell = sheet.cell(row, column)
            x1, x2 = x_positions[column - 1], x_positions[end_column]
            y_top, y_bottom = y_positions[row - 1], y_positions[end_row]
            width, height = x2 - x1, y_top - y_bottom
            fill = _fill_rgb(cell)
            if fill:
                document.setFillColorRGB(*fill)
                document.rect(x1, y_bottom, width, height, fill=1, stroke=0)

            edge_cells = {
                "left": sheet.cell(row, column),
                "right": sheet.cell(row, end_column),
                "top": sheet.cell(row, column),
                "bottom": sheet.cell(end_row, column),
            }
            document.setStrokeColorRGB(0, 0, 0)
            document.setLineWidth(0.55)
            if edge_cells["top"].border.top.style:
                document.line(x1, y_top, x2, y_top)
            if edge_cells["bottom"].border.bottom.style:
                document.line(x1, y_bottom, x2, y_bottom)
            if edge_cells["left"].border.left.style:
                document.line(x1, y_bottom, x1, y_top)
            if edge_cells["right"].border.right.style:
                document.line(x2, y_bottom, x2, y_top)

            value = "" if cell.value is None else str(cell.value)
            if not value or value in {"0", "0.0", ","}:
                continue
            font_name = "AutoDogovorBold" if cell.font.bold else "AutoDogovor"
            font_size = max(6.8, min(13.5, float(cell.font.sz or 10) * 1.03))
            text_width = max(4, width - 5)
            lines = _wrap_pdf_text(value, font_name, font_size, text_width)
            while lines and len(lines) * font_size * 1.08 > height - 2 and font_size > 5.8:
                font_size -= 0.35
                lines = _wrap_pdf_text(value, font_name, font_size, text_width)
            line_height = font_size * 1.08
            total_height = len(lines) * line_height
            vertical = cell.alignment.vertical or "bottom"
            if vertical == "top":
                baseline = y_top - font_size - 1.5
            elif vertical == "center":
                baseline = y_bottom + (height + total_height) / 2 - line_height
            else:
                baseline = y_bottom + total_height - line_height + 1.5
            document.setFillColorRGB(0.04, 0.04, 0.04)
            document.setFont(font_name, font_size)
            for line in lines:
                line_width = pdfmetrics.stringWidth(line, font_name, font_size)
                horizontal = cell.alignment.horizontal or "left"
                if horizontal in {"center", "centerContinuous"}:
                    text_x = x1 + (width - line_width) / 2
                elif horizontal == "right":
                    text_x = x2 - line_width - 2.5
                else:
                    text_x = x1 + 2.5
                document.drawString(text_x, baseline, line)
                baseline -= line_height
    document.showPage()
    document.save()
    workbook.close()


def _render_xls_pdf(xls_path: Path, pdf_path: Path) -> None:
    """Render the legacy template with its native row/column geometry to one A4 page."""
    _register_pdf_fonts()
    workbook = xlrd.open_workbook(str(xls_path), formatting_info=True)
    sheet = workbook.sheet_by_index(0)
    page_width, page_height = A4
    margin_x, margin_y = 24.0, 22.0
    column_units = [(sheet.colinfo_map.get(column).width / 256 if sheet.colinfo_map.get(column) else 8.43) for column in range(10)]
    row_units = [(sheet.rowinfo_map.get(row).height / 20 if sheet.rowinfo_map.get(row) else 15) for row in range(47)]
    width_scale = (page_width - margin_x * 2) / sum(column_units)
    height_scale = (page_height - margin_y * 2) / sum(row_units)
    x_positions = [margin_x]
    for width in column_units:
        x_positions.append(x_positions[-1] + width * width_scale)
    y_positions = [page_height - margin_y]
    for height in row_units:
        y_positions.append(y_positions[-1] - height * height_scale)

    anchors: dict[tuple[int, int], tuple[int, int]] = {}
    covered: set[tuple[int, int]] = set()
    for row_start, row_end, column_start, column_end in sheet.merged_cells:
        anchors[(row_start, column_start)] = (row_end - 1, column_end - 1)
        for row in range(row_start, row_end):
            for column in range(column_start, column_end):
                if (row, column) != (row_start, column_start):
                    covered.add((row, column))

    document = canvas.Canvas(str(pdf_path), pagesize=A4, pageCompression=1)
    for row in range(47):
        for column in range(10):
            if (row, column) in covered:
                continue
            end_row, end_column = anchors.get((row, column), (row, column))
            cell = sheet.cell(row, column)
            xf = workbook.xf_list[cell.xf_index]
            font = workbook.font_list[xf.font_index]
            border = xf.border
            alignment = xf.alignment
            x1, x2 = x_positions[column], x_positions[end_column + 1]
            y_top, y_bottom = y_positions[row], y_positions[end_row + 1]
            document.setStrokeColorRGB(0, 0, 0)
            document.setLineWidth(.45)
            if border.top_line_style: document.line(x1, y_top, x2, y_top)
            if border.bottom_line_style: document.line(x1, y_bottom, x2, y_bottom)
            if border.left_line_style: document.line(x1, y_bottom, x1, y_top)
            if border.right_line_style: document.line(x2, y_bottom, x2, y_top)
            value = "" if cell.value is None else str(cell.value)
            if not value or value in {"0", "0.0", ","}:
                continue
            font_name = "AutoDogovorBold" if font.bold else "AutoDogovor"
            font_size = max(5.8, min(13.5, float(font.height or 200) / 20 * .98))
            lines = _wrap_pdf_text(value, font_name, font_size, max(4, x2 - x1 - 5))
            while lines and len(lines) * font_size * 1.08 > y_top - y_bottom - 2 and font_size > 5.5:
                font_size -= .35
                lines = _wrap_pdf_text(value, font_name, font_size, max(4, x2 - x1 - 5))
            line_height = font_size * 1.08
            baseline = y_bottom + (y_top - y_bottom + len(lines) * line_height) / 2 - line_height
            document.setFillColorRGB(.04, .04, .04)
            document.setFont(font_name, font_size)
            for line in lines:
                line_width = pdfmetrics.stringWidth(line, font_name, font_size)
                if alignment.hor_align == 2:
                    text_x = x1 + (x2 - x1 - line_width) / 2
                elif alignment.hor_align == 3:
                    text_x = x2 - line_width - 2.5
                else:
                    text_x = x1 + 2.5
                document.drawString(text_x, baseline, line)
                baseline -= line_height
    document.showPage()
    document.save()


def _render_pdf_jpg(pdf_path: Path, jpg_path: Path) -> None:
    document = pdfium.PdfDocument(str(pdf_path))
    if len(document) != 1:
        raise ValueError("Договор должен помещаться на одной странице")
    page = document[0]
    bitmap = page.render(scale=300 / 72)
    image: Image.Image = bitmap.to_pil().convert("RGB").resize((2480, 3508))
    image.save(jpg_path, "JPEG", quality=95, optimize=True)
    bitmap.close()
    page.close()
    document.close()


def create_contract(deal_id: int, payload: dict[str, Any]) -> list[dict[str, str]]:
    xlsx_path = _output_path(deal_id, payload, ".xls")
    pdf_path = _output_path(deal_id, payload, ".pdf")
    jpg_path = _output_path(deal_id, payload, ".jpeg")
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError("Не найден шаблон MyFiles/BAZA.xls")
    _fill_legacy_template(xlsx_path, payload)
    _render_xls_pdf(xlsx_path, pdf_path)
    _render_pdf_jpg(pdf_path, jpg_path)
    return [
        {
            "kind": "xls",
            "name": xlsx_path.name,
            "stored_path": str(xlsx_path),
            "media_type": "application/vnd.ms-excel",
        },
        {"kind": "pdf", "name": pdf_path.name, "stored_path": str(pdf_path), "media_type": "application/pdf"},
        {"kind": "jpg", "name": jpg_path.name, "stored_path": str(jpg_path), "media_type": "image/jpeg"},
    ]
